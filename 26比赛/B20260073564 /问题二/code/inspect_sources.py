"""Extract the supplied DOCX (including math runs) and workbook cells for audit."""
from pathlib import Path
import csv
import json
import zipfile
from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT.parent / 'B题'
OUT = ROOT / 'inputs'
OUT.mkdir(exist_ok=True, parents=True)
doc = BASE / '氢燃料电池低温冷启动建模与控制策略研究.docx'
with zipfile.ZipFile(doc) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
lines = []
for p in xml.xpath('//w:body//w:p', namespaces=ns):
    text = ''.join(p.xpath('.//w:t/text() | .//m:t/text()', namespaces=ns))
    if text.strip():
        lines.append(text)
(OUT / '题目全文提取.txt').write_text('\n'.join(lines), encoding='utf-8')
summary = {}
for i in (1, 2):
    source = BASE / '氢燃料电池低温冷启动建模与控制策略研究  附件' / f'附件{i}.xlsx'
    wb = load_workbook(source, data_only=True, read_only=True)
    for ws in wb:
        rows = list(ws.values)
        name = f'附件{i}_{ws.title}'
        with (OUT / f'{name}.csv').open('w', newline='', encoding='utf-8-sig') as f:
            csv.writer(f).writerows(rows)
        summary[name] = {'rows': ws.max_row, 'columns': ws.max_column, 'preview': rows[:6]}
print(json.dumps(summary, ensure_ascii=False, indent=2))
print('\n'.join(lines))

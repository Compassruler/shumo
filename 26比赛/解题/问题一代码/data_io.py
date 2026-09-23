"""附件只读输入与结果表输出；电流密度严格读取附件2 F列（已经是SI）。"""
from pathlib import Path
import csv
import hashlib
import sys

HERE = Path(__file__).resolve().parent
# 此目录中的二进制依赖是Windows CPython 3.12构建；其他环境使用自行安装的requirements。
if sys.platform == "win32" and sys.version_info[:2] == (3, 12) and (HERE / ".python_deps").exists():
    sys.path.insert(0, str(HERE / ".python_deps"))
sys.path.insert(0, str(HERE.parent))
import numpy as np
from openpyxl import load_workbook


def attachment_paths():
    folder = HERE.parent / "B题" / "氢燃料电池低温冷启动建模与控制策略研究  附件"
    return folder / "附件1.xlsx", folder / "附件2.xlsx"


def read_experiments(path=None):
    """F列缓存缺失时直接报错，绝不静默退回E列或用I/25重算。"""
    path = Path(path) if path else attachment_paths()[1]
    book = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    result = {}
    for initial in (-20, -25):
        name = f"{initial}℃"
        sheet = book[name]
        header = str(sheet["F2"].value)
        if "电流密度" not in header:
            raise ValueError(f"{name}!F2不是电流密度标签：{header}")
        rows, excel_rows = [], []
        for row_number, row in enumerate(sheet.iter_rows(min_row=3, max_col=6,
                                                        values_only=True), 3):
            if all(v is None for v in row):
                continue
            if any(not isinstance(v, (int, float)) for v in row):
                raise ValueError(f"{name}!A{row_number}:F{row_number}含空值或非数值；"
                                 "请在Excel中重算并保存公式缓存。")
            rows.append(row)
            excel_rows.append(row_number)
        a = np.asarray(rows, dtype=float)
        if not len(a) or not np.all(np.isfinite(a)):
            raise ValueError(f"{name}无有效数值数据")
        if np.any(np.diff(a[:, 0]) <= 0) or a[0, 0] != 0:
            raise ValueError(f"{name}时间必须从0开始且严格递增")
        if np.any(a[:, 5] < 0):
            raise ValueError(f"{name}不支持负电流密度")
        area = a[a[:, 5] > 0, 1] / a[a[:, 5] > 0, 5] * 1e4
        result[initial] = dict(
            initial_celsius=initial, sheet=name, t=a[:, 0], current=a[:, 5],
            voltage_exp=a[:, 2], temperature_exp=a[:, 3],
            audit=dict(source=str(path), sheet=name, input_column="F", unit="A/m²",
                       count=len(a), first_row=excel_rows[0], last_row=excel_rows[-1],
                       time_range_s=[float(a[0, 0]), float(a[-1, 0])],
                       current_range_A_m2=[float(a[:, 5].min()), float(a[:, 5].max())],
                       F3_formula=formulas[name]["F3"].value,
                       F_equals_E_times_10000=bool(np.allclose(a[:, 5], a[:, 4]*1e4)),
                       implied_area_cm2_min=float(area.min()),
                       implied_area_cm2_median=float(np.median(area)),
                       implied_area_cm2_max=float(area.max()),
                       attachment1_area_cm2=25.0,
                       sha256=hashlib.sha256(path.read_bytes()).hexdigest()),
        )
    book.close()
    formulas.close()
    return result


def write_csv(path, columns):
    """UTF-8 BOM便于Windows Excel直接打开，原始附件不修改。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns.keys())
        writer.writerows(zip(*columns.values()))


def error_metrics(simulation, experiment):
    d = np.asarray(simulation) - np.asarray(experiment)
    denom = np.abs(experiment)
    rel = np.divide(np.abs(d)*100, denom, out=np.full_like(d, np.nan), where=denom > 1e-12)
    return dict(mae=float(np.mean(np.abs(d))), rmse=float(np.sqrt(np.mean(d*d))),
                max_abs=float(np.max(np.abs(d))), mean_relative_percent=float(np.nanmean(rel)),
                max_relative_percent=float(np.nanmax(rel)))


if __name__ == "__main__":
    for temperature, data in read_experiments().items():
        print(temperature, data["audit"])
